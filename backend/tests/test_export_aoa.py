# -*- coding: utf-8 -*-
"""`POST /api/factors/export/aoa-xlsx`（★ v2.11.68）—— **通用二维表导出**。

## 为什么加它（安全债 P2-#22 的正解）
前端原来用 npm `xlsx@0.18.5`（SheetJS）**在前端写 Excel** ✗，那个包已停止维护
⇒ 两条 high **无 fix**（原型链污染 / ReDoS）✓。
试过换 `exceljs`：**更差** ✗（185 个包、审计变 3 漏洞含 1 high ✓）。
⇒ 回到本项目既有约定：**导出由后端 openpyxl 生成** ✓（`exports/xl.py` 一直在用）⇒ 零新依赖 ✓。

## 这三条断言为什么重要
① **数字必须还是数字** ✗→✓ —— 全变文本的话 Excel 里**没法排序/求和** ✓（用户当场会发现 ✗）；
② **页名要清洗** —— Excel 不许 `[]:*?/\\`，openpyxl 会**直接抛** ⇒ 用户看到 500 ✗；
③ **必须有上限** —— 无界导出 = 一次 POST 就能把服务撑爆 ✗。
"""
from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)
URL = "/api/factors/export/aoa-xlsx"


def _post(sheets, filename="测试导出.xlsx"):
    return client.post(URL, json={"filename": filename, "sheets": sheets})


def _read(resp):
    assert resp.status_code == 200, resp.text[:300]
    return load_workbook(BytesIO(resp.content))


def test_two_sheets_roundtrip():
    r = _post([
        {"name": "因子指标", "rows": [["因子", "IC", "结论"], ["mom20", 0.031, "有效"]]},
        {"name": "因子与公式", "rows": [["因子", "来源", "公式"], ["mom20", "内置", "Mean($close,20)"]]},
    ])
    wb = _read(r)
    assert wb.sheetnames == ["因子指标", "因子与公式"]
    ws = wb["因子指标"]
    assert ws.cell(1, 1).value == "因子"
    assert ws.cell(2, 1).value == "mom20"


def test_numbers_stay_numbers():
    """⭐⭐ 数字若被写成文本，Excel 里**没法排序/求和** ✗ —— 这是导出最容易踩的坑 ✓。"""
    r = _post([{"name": "s", "rows": [["a", "b"], [1.25, -3], [0, 7]]}])
    ws = _read(r)["s"]
    assert isinstance(ws.cell(2, 1).value, (int, float)), "1.25 被写成了 %r ✗" % (ws.cell(2, 1).value,)
    assert abs(float(ws.cell(2, 1).value) - 1.25) < 1e-9
    assert int(ws.cell(2, 2).value) == -3
    assert int(ws.cell(3, 2).value) == 7


def test_empty_cells_stay_empty():
    """`None` / 空串 ⇒ **空单元格**（不是 0 ✗ —— 0 会被当成真实数值算进均值 ✓）。"""
    r = _post([{"name": "s", "rows": [["a", "b", "c"], [None, "", "x"]]}])
    ws = _read(r)["s"]
    assert ws.cell(2, 1).value is None
    assert ws.cell(2, 2).value is None
    assert ws.cell(2, 3).value == "x"


def test_illegal_sheet_name_is_sanitized():
    """⭐ 页名含 `[]:*?/\\` ⇒ openpyxl **会直接抛** ⇒ 用户看到 500 ✗ ⇒ 必须清洗 ✓。"""
    r = _post([{"name": "a/b:c*d?", "rows": [["x"]]}])
    wb = _read(r)
    assert len(wb.sheetnames) == 1
    for ch in "[]:*?/\\":
        assert ch not in wb.sheetnames[0], "页名没清洗干净：%r" % (wb.sheetnames[0],)
    assert wb.sheetnames[0].startswith("a"), wb.sheetnames[0]


def test_long_sheet_name_is_truncated():
    r = _post([{"name": "页" * 50, "rows": [["x"]]}])
    assert len(_read(r).sheetnames[0]) <= 31


def test_filename_header_is_returned():
    r = _post([{"name": "s", "rows": [["x"]]}], filename="单因子测试结果.xlsx")
    assert r.status_code == 200
    assert "X-Export-Filename" in r.headers, "前端靠这个头显示文件名 ✓（CORS 已 expose ✓）"
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def test_empty_payload_is_400():
    r = client.post(URL, json={"filename": "x.xlsx", "sheets": []})
    assert r.status_code == 400
    r2 = client.post(URL, json={"filename": "x.xlsx", "sheets": [{"name": "s", "rows": []}]})
    assert r2.status_code == 400, "全是空页 ⇒ 也该 400（否则生成一个空文件 ✗）"


def test_too_many_sheets_is_400():
    sheets = [{"name": "s%d" % i, "rows": [["x"]]} for i in range(25)]
    assert _post(sheets).status_code == 400
