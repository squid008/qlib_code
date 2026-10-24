# -*- coding: utf-8 -*-
"""`exports/exit_curve.build_exit_curve_workbook`（★ v2.11.75）—— 自包含的 XLSX 组装。

## 为什么值得测
它是**用户拿到手的那份文件** ✓，坏掉的方式都很"体面" ✗：
· **页序错** ⇒ 用户第一眼看到的是「离场A」而不是曲线截图（而这文件的第一页**必须**是截图 ✓）；
· **没带截图时不留说明** ⇒ 用户以为"图丢了"✗，其实是浏览器抓图失败 ✓（必须写清"看下面哪几节"✓）；
· **调仓表头少一列** ⇒ 后面的数据整体错位 ✗（而 Excel 里"看起来"只是少了"原因"列 ✓）；
· **失败时出半本书** ⇒ 用户拿到一个能打开但缺页的文件 ✗（本函数的约定是**失败向上抛** ✓）。
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.exports.exit_curve import (
    SHEET_SHOT,
    TRADE_HEADERS,
    ExitCurveDoc,
    ExitRunDoc,
    build_exit_curve_workbook,
)

#: 1×1 透明 PNG（合法文件 ✓ —— 只为走通"带截图"那条分支）
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
    "890000000a49444154789c6300010000050001 0d0a2db40000000049454e44ae426082".replace(" ", "")
)


def _texts(ws) -> str:
    """把一个 sheet 的所有单元格文本拼起来（断言"文案在不在"用 ✓）。"""
    out = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out.append(str(c.value))
    return "\n".join(out)


def _build(doc: ExitCurveDoc):
    data = build_exit_curve_workbook(doc)
    assert isinstance(data, bytes) and data[:2] == b"PK", "产物必须是合法 xlsx（PK 魔数 ✓）"
    return load_workbook(BytesIO(data))


class TestSheets:
    def test_first_sheet_is_always_the_shot(self):
        """⭐⭐ 第 1 页**必须**是截图页 ✓（用户指定；页序 = 创建顺序 ✓）。"""
        doc = ExitCurveDoc(runs=[ExitRunDoc(sheet="离场A", title="A"), ExitRunDoc(sheet="离场B")])
        wb = _build(doc)
        assert wb.sheetnames[0] == SHEET_SHOT, wb.sheetnames
        assert wb.sheetnames[1:] == ["离场A", "离场B"]

    def test_one_page_per_run_in_order(self):
        runs = [ExitRunDoc(sheet="s%d" % i) for i in range(4)]
        assert _build(ExitCurveDoc(runs=runs)).sheetnames[1:] == ["s0", "s1", "s2", "s3"]

    def test_no_runs_still_valid(self):
        wb = _build(ExitCurveDoc())
        assert wb.sheetnames == [SHEET_SHOT]


class TestShotPage:
    def test_missing_shot_has_explanation(self):
        """⭐ 没带截图时必须**说明原因和替代读法** ✓ —— 否则用户以为"图丢了"✗。"""
        t = _texts(_build(ExitCurveDoc()).sheetnames and _build(ExitCurveDoc())[SHEET_SHOT])
        assert "没带上曲线截图" in t
        assert "参数与口径" in t, "要指路（看哪一节）✓"
        assert "调仓记录" in t

    def test_with_shot_does_not_print_missing_note(self):
        t = _texts(_build(ExitCurveDoc(shot_png=PNG_1PX))[SHEET_SHOT])
        assert "没带上曲线截图" not in t

    def test_title_and_headline(self):
        doc = ExitCurveDoc(title="离场曲线：动量", headline=["基准 SH000300", ""],
                           subtitle="说明一句")
        t = _texts(_build(doc)[SHEET_SHOT])
        assert "离场曲线：动量" in t and "基准 SH000300" in t and "说明一句" in t

    def test_params_section_only_when_present(self):
        # ⚠ `param_pairs` 是**3 元组**（`kv(per_row=1)` 要 label/value/说明 三列 ✓）——
        #   我第一版给了 2 元组 ✗ ⇒ `ValueError: not enough values to unpack (expected 3)`
        with_p = _texts(_build(ExitCurveDoc(
            param_pairs=[("K", "10", ""), ("持仓", "20天", "交易日")]))[SHEET_SHOT])
        assert "二、参数与口径" in with_p and "K" in with_p
        without = _texts(_build(ExitCurveDoc())[SHEET_SHOT])
        assert "二、参数与口径" not in without, "没有参数就不该留一个空分区 ✗"

    def test_page_list_and_notes(self):
        doc = ExitCurveDoc(runs=[ExitRunDoc(sheet="离场A", title="A", headline=["年化 12%"])],
                           notes=["备注一", "备注二"])
        t = _texts(_build(doc)[SHEET_SHOT])
        assert "三、本文件包含的页" in t and "离场A" in t and "年化 12%" in t
        assert "备注一" in t and "备注二" in t


class TestRunPages:
    @staticmethod
    def _header_row(ws) -> int:
        """找表头在第几行 —— ⚠ 不在第 1 行 ✗（上面先有标题/headline ✓，
        我第一版直接读 row 1 ⇒ 拿到的是**标题**，当场红 ✓）。"""
        for r in range(1, ws.max_row + 1):
            if ws.cell(r, 1).value == TRADE_HEADERS[0]:
                return r
        raise AssertionError("没找到调仓表头（首列应是「日期」）✗")

    def test_trade_table_headers_exact(self):
        """⭐⭐ 表头必须**完整且有序** ✓ —— 少一列会让后面所有数据整体错位 ✗。"""
        trades = [["2024-01-02", "SH600000", "卖", 100, 10.5, 1050.0, 0.5, "exit_signal"]]
        wb = _build(ExitCurveDoc(runs=[ExitRunDoc(sheet="A", title="A", trades=trades)]))
        ws = wb["A"]
        r0 = self._header_row(ws)
        row = [ws.cell(r0, i + 1).value for i in range(len(TRADE_HEADERS))]
        assert row == TRADE_HEADERS, "表头变了 ✗：%s" % row

    def test_trades_rows_written(self):
        trades = [["2024-01-02", "SH600000", "卖", 100, 10.5, 1050.0, 0.5, "exit_signal"],
                  ["2024-01-03", "SZ000001", "买", 200, 8.0, 1600.0, 0.8, "rebalance"]]
        wb = _build(ExitCurveDoc(runs=[ExitRunDoc(sheet="A", trades=trades)]))
        ws = wb["A"]
        r0 = self._header_row(ws)
        assert ws.cell(r0 + 1, 1).value == "2024-01-02"
        assert ws.cell(r0 + 2, 2).value == "SZ000001"

    def test_empty_trades_skips_table(self):
        t = _texts(_build(ExitCurveDoc(runs=[ExitRunDoc(sheet="A", title="A")]))["A"])
        assert TRADE_HEADERS[0] not in t or "日期" not in t.split("\n")[1:], "没成交就不该出空表 ✗"

    def test_run_headline_and_title(self):
        doc = ExitCurveDoc(runs=[ExitRunDoc(sheet="A", title="离场A：止损", headline=["年化 12%", ""])])
        t = _texts(_build(doc)["A"])
        assert "离场A：止损" in t and "年化 12%" in t


def test_returns_bytes_not_path(tmp_path):
    """约定：**返回 bytes**（内存里生成 ✓）—— 落临时文件会在并发下互相踩 ✗。"""
    before = set(p.name for p in tmp_path.iterdir())
    build_exit_curve_workbook(ExitCurveDoc(runs=[ExitRunDoc(sheet="A")]))
    assert set(p.name for p in tmp_path.iterdir()) == before


def test_raises_on_bad_doc():
    """⭐ 失败要**向上抛** ✓（路由转成可读 500 ✓）—— 绝不静默出半本书 ✗。"""

    class _Bad:
        runs = None            # 迭代它会 TypeError ✓
        title = "x"
        headline: list = []
        subtitle = ""
        shot_png = None
        shot_title = ""
        shot_note = ""
        param_pairs: list = []
        notes: list = []

    # ⚠ 这里被**我自己加的 ruff `B017`** 抓到一次 ✗ —— 原来写的是 `pytest.raises(Exception)`
    #   ⇒ 等于"只要炸了就算对" ✓（连实现 bug 也算预期 ✗）。具体是 `for run in None` ⇒ TypeError ✓。
    with pytest.raises(TypeError):
        build_exit_curve_workbook(_Bad())  # type: ignore[arg-type]
