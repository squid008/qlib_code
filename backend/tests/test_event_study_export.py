# -*- coding: utf-8 -*-
"""「事件研究」XLSX 导出的契约守卫（v2.1.0 ✓）。

**为什么锁**：事件研究的导出与单因子测试**共用同一套回灌契约**（`hold_days`/`cost`/`benchmark`
三个键同名同义 ✓）⇒ 一旦这里写歪，"拖回交易信号测试复现净值"就会静默取错参数 ✗。
另外第 2 页的持仓列也是契约（`HOLDING_COLS[KIND_EVENT_STUDY]` ✓）⇒ 列名/顺序必须与
`positions_from_trades()` 的字段一一对上 ✓。
"""
from io import BytesIO

from openpyxl import load_workbook

from app.exports import spec
from app.exports.event_study import (EventStudyExportDoc, build_event_study_workbook,
                                      contract_params, param_pairs_from)


def _png(w: int = 200, h: int = 80) -> bytes:
    from PIL import Image

    buf = BytesIO()
    Image.new("RGB", (w, h), "#0f766e").save(buf, format="PNG")
    return buf.getvalue()


def _doc() -> EventStudyExportDoc:
    params = contract_params(
        factor_name="AR13_ID60", factor_expression="$close>Ref($close,1)",
        universe="all", start="2022-01-01", end="2026-08-21", max_k=60,
        hold_days=60, cost=0.004, capital=1e9, benchmark="SH000300",
        app_version="9.9.9", generated_at="2026-10-11 12:00:00",
    )
    return EventStudyExportDoc(
        factor_name="AR13_ID60",
        headline="all · 2022-01-01~2026-08-21 · 最长持有 60 日 · 触发 3210 次（纳入 3000）",
        subtitle="导出时间 2026-10-11 12:00:00",
        params=params,
        formula_pairs=[("因子名称", "AR13_ID60", None)],
        param_pairs=param_pairs_from(params, skip=(spec.P_FACTOR_NAME, spec.P_FACTOR_EXPR)),
        kpi=[("触发事件数", 3210, "0"), ("持有 60 日中位数", 0.031, "0.0%")],
        kpi_png=_png(),
        sections=[("持有期收益曲线", "T+1 收盘买入；单位 %", _png()),
                  ("净值曲线（两种资金方案 + 基准）", "持有周期 60 交易日", _png())],
        notes=["· 两种资金方案：event_even / cash_even ✓"],
        holdings=[{"rebalance_date": "2022-01-04", "mode": "event_even", "code": "SH600001",
                   "shares": 100, "price": 12.34, "amount": 1234.0, "weight": 0.5},
                  {"rebalance_date": "2022-01-04", "mode": "cash_even", "code": "SH600002",
                   "shares": 200, "price": 6.17, "amount": 1234.0, "weight": 1.0}],
        holdings_note="· 第 2 页持仓由成交明细派生 ✓",
    )


def test_event_study_workbook_and_contract():
    doc = _doc()
    data = build_event_study_workbook(doc)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames[:3] == [spec.SHEET_PARAMS, spec.SHEET_HOLDINGS, spec.SHEET_NOTES]
    assert spec.SHEET_META in wb.sheetnames and wb[spec.SHEET_META].sheet_state == "hidden"
    # 摘要卡片 + 两个图块 = 3 张截图（用户 2026-10-11 要求"触发事件统计也用截图的形式" ✓）
    assert len(wb[spec.SHEET_PARAMS]._images) == 3
    assert wb[spec.SHEET_PARAMS].freeze_panes is None       # 第 1 页**不能**冻结（会滚不动 ✗）
    assert wb[spec.SHEET_HOLDINGS].freeze_panes == "A2"     # 持仓页要冻表头 ✓

    out = spec.read_export(data)
    assert out.kind == spec.KIND_EVENT_STUDY
    # ★ 回灌三件套（与单因子测试导出**同键同义** ✓）
    assert out.p(spec.P_HOLD_DAYS) == 60 and abs(float(out.p(spec.P_COST)) - 0.004) < 1e-12
    assert out.p(spec.P_BENCHMARK) == "SH000300" and out.p(spec.P_MAX_K) == 60
    # 持仓列 = 契约（代码/股数/成交价/金额/权重 ✓）
    assert out.holdings and set(out.holdings[0]) >= {"rebalance_date", "mode", "code", "shares",
                                                     "price", "amount", "weight"}
    assert out.holdings[0]["code"] == "SH600001" and out.n_periods() == 1
    # 列的**中文表头**必须与契约一致（回灌按表头反查 ✓）
    assert set(out.holding_labels) == {zh for _k, zh in spec.HOLDING_COLS[spec.KIND_EVENT_STUDY]}
